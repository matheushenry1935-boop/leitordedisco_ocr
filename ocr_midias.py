"""
OCR MÍDIAS JUDICIAIS V6 - leitura por IA (visão) - Windows.

Fluxo:
  1. Câmera ao vivo  ->  Espaço captura a foto
  2. A IA lê o disco (nº do processo, data escrita no CD, o que ocorreu)
  3. O formulário é preenchido para você conferir/editar
  4. Ctrl+S (ou "Próximo CD") grava o item e esvazia o formulário
  5. Botão verde: pergunta o nome do lote (informações do envelope) e exporta o Excel

Atalhos: Espaço = capturar | Ctrl+S = Próximo CD (salvar item)
A chave de API fica em config.ini (botão "Configurar IA").
"""
import base64
import configparser
import json
import os
import queue
import re
import sqlite3
import threading
import time
import tkinter as tk
import urllib.error
import urllib.request
from datetime import datetime
from pathlib import Path
from tkinter import messagebox, simpledialog

import cv2
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from PIL import Image, ImageTk

BASE = Path(__file__).resolve().parent
DB_PATH = BASE / "levantamento_midias.db"
EXPORT_DIR = BASE / "exportacoes"
CONFIG_PATH = BASE / "config.ini"

# Tema escuro
BG, PAINEL, CAMPO = "#1b1c26", "#262838", "#12131a"
TXT, SUB, CIANO = "#e8e9f0", "#aeb1c4", "#22d3ee"
AZUL, VERDE, BORDA = "#2563eb", "#16a34a", "#3a3d52"
OK_COR, AVISO_COR = "#22c55e", "#f59e0b"
FONTE = "Segoe UI"
PREV_W, PREV_H = 260, 140

# --------------------------------------------------------------------------
# Provedores de IA (todos com leitura de imagem). Troque o modelo em "Configurar IA".
# --------------------------------------------------------------------------
PROVEDORES = {
    "anthropic": {"nome": "Anthropic (Claude)", "modelo": "claude-sonnet-5-5",
                  "url": "https://api.anthropic.com/v1/messages", "env": "ANTHROPIC_API_KEY"},
    "gemini": {"nome": "Google (Gemini)", "modelo": "gemini-3.6-flash",
               "url": "https://generativelanguage.googleapis.com/v1beta/models/{modelo}:generateContent",
               "env": "GEMINI_API_KEY"},
    "openai": {"nome": "OpenAI (ChatGPT)", "modelo": "gpt-5.3-chat-latest",
               "url": "https://api.openai.com/v1/chat/completions", "env": "OPENAI_API_KEY"},
}

PROMPT = """Você está ajudando a catalogar mídias (CDs/DVDs) de audiências da Justiça Federal brasileira.
A imagem mostra um disco ou sua capa, com informações escritas à mão ou impressas. Leia com muita atenção.
Responda SOMENTE com um objeto JSON (sem markdown, sem texto fora do JSON) com estas chaves:

"numero_processo": número do processo judicial no formato CNJ NNNNNNN-DD.AAAA.J.TR.OOOO (20 dígitos).
   Se houver mais de um número, use o que está junto de "Autos"/"Processo" e cite os demais em "observacoes".
   Copie os dígitos exatamente como estão visíveis. NUNCA invente dígitos; se estiver ilegível, use "".
"data": data da audiência/ato ESCRITA NO DISCO, no formato dd/mm/aaaa (dia/mês/ano, padrão brasileiro).
   Se o ano tiver 2 dígitos, converta para 4 (ex.: 18 vira 2018). Se não houver data escrita, use "".
   NUNCA use a data de hoje.
"hora": horário escrito no disco no formato HH:MM, ou "".
"ocorrido": o que ocorreu / tipo do ato, em poucas palavras, conforme escrito (ex.: "Audiência de instrução",
   "Interrogatório", "Oitiva de testemunhas"). Se não houver, "".
"materia": "Criminal" ou "Cível" SOMENTE se o texto indicar claramente (ex.: ação penal, réu, denúncia,
   inquérito = Criminal). Caso contrário, "A Verificar".
"transcricao": todo o texto legível no disco, linha por linha, separado por \\n, como aparece.
"observacoes": anotações relevantes que não couberam acima (outros números, nomes, trechos ilegíveis).
   Use "" se não houver.

Se uma informação não estiver visível, deixe "". Não adivinhe."""


# --------------------------------------------------------------------------
# Configuração
# --------------------------------------------------------------------------
def carregar_config() -> dict:
    cp = configparser.ConfigParser(interpolation=None)
    cp.read(CONFIG_PATH, encoding="utf-8")
    sec = cp["ia"] if cp.has_section("ia") else {}
    prov = sec.get("provedor", "anthropic")
    if prov not in PROVEDORES:
        prov = "anthropic"
    chave = sec.get("chave", "").strip() or os.environ.get(PROVEDORES[prov]["env"], "").strip()
    modelo = sec.get("modelo", "").strip() or PROVEDORES[prov]["modelo"]
    return {"provedor": prov, "chave": chave, "modelo": modelo}


def salvar_config(cfg: dict):
    cp = configparser.ConfigParser(interpolation=None)
    cp["ia"] = {k: cfg[k] for k in ("provedor", "chave", "modelo")}
    with open(CONFIG_PATH, "w", encoding="utf-8") as f:
        cp.write(f)


# --------------------------------------------------------------------------
# Leitura por IA
# --------------------------------------------------------------------------
def foto_para_base64(bgr, lado_max=1456) -> str:
    """Reduz a foto (mais barato e rápido) e converte para JPEG em base64."""
    h, w = bgr.shape[:2]
    esc = lado_max / max(h, w)
    if esc < 1:
        bgr = cv2.resize(bgr, (int(w * esc), int(h * esc)), interpolation=cv2.INTER_AREA)
    ok, buf = cv2.imencode(".jpg", bgr, [cv2.IMWRITE_JPEG_QUALITY, 90])
    if not ok:
        raise RuntimeError("Falha ao preparar a imagem.")
    return base64.b64encode(buf.tobytes()).decode("ascii")


def _msg_erro(e: urllib.error.HTTPError) -> str:
    try:
        j = json.loads(e.read().decode("utf-8", "replace"))
        err = j.get("error", j)
        msg = err.get("message") if isinstance(err, dict) else str(err)
        return (msg or str(e.reason))[:250]
    except Exception:
        return str(e.reason)[:250]


def _post_json(url, headers, corpo, tentativas=3, timeout=60) -> dict:
    """POST com nova tentativa em falhas temporárias (limite de uso, servidor ocupado, rede)."""
    dados = json.dumps(corpo).encode("utf-8")
    for i in range(tentativas):
        ultima = i == tentativas - 1
        try:
            req = urllib.request.Request(url, data=dados, headers=headers, method="POST")
            with urllib.request.urlopen(req, timeout=timeout) as r:
                return json.loads(r.read().decode("utf-8"))
        except urllib.error.HTTPError as e:  # precisa vir antes de URLError/OSError
            if e.code in (408, 429, 500, 502, 503, 504, 529) and not ultima:
                time.sleep(2 * (i + 1))
                continue
            dica = " Verifique a chave de API." if e.code in (401, 403) else ""
            raise RuntimeError(f"Erro {e.code} da IA: {_msg_erro(e)}.{dica}") from None
        except OSError:
            if not ultima:
                time.sleep(2 * (i + 1))
                continue
            raise RuntimeError("Sem conexão com a IA (verifique a internet ou o proxy).") from None


def chamar_ia(cfg: dict, b64: str, prompt: str) -> str:
    """Envia imagem + instruções ao provedor escolhido e devolve o texto da resposta."""
    prov, chave, modelo = cfg["provedor"], cfg["chave"], cfg["modelo"]
    url = PROVEDORES[prov]["url"].format(modelo=modelo)
    if prov == "anthropic":
        headers = {"x-api-key": chave, "anthropic-version": "2023-06-01", "content-type": "application/json"}
        corpo = {"model": modelo, "max_tokens": 4000, "messages": [{"role": "user", "content": [
            {"type": "image", "source": {"type": "base64", "media_type": "image/jpeg", "data": b64}},
            {"type": "text", "text": prompt}]}]}
        r = _post_json(url, headers, corpo)
        texto = "".join(b.get("text", "") for b in r.get("content", []) if b.get("type") == "text")
    elif prov == "gemini":
        headers = {"x-goog-api-key": chave, "content-type": "application/json"}
        corpo = {"contents": [{"parts": [
            {"inline_data": {"mime_type": "image/jpeg", "data": b64}}, {"text": prompt}]}],
            "generationConfig": {"responseMimeType": "application/json"}}
        r = _post_json(url, headers, corpo)
        cands = r.get("candidates") or []
        if not cands:
            motivo = (r.get("promptFeedback") or {}).get("blockReason", "sem candidatos")
            raise RuntimeError(f"A IA não devolveu resposta ({motivo}).")
        partes = (cands[0].get("content") or {}).get("parts") or []
        texto = "".join(p.get("text", "") for p in partes if not p.get("thought"))
    else:  # openai
        headers = {"Authorization": f"Bearer {chave}", "content-type": "application/json"}
        corpo = {"model": modelo, "response_format": {"type": "json_object"}, "messages": [{"role": "user", "content": [
            {"type": "text", "text": prompt},
            {"type": "image_url", "image_url": {"url": "data:image/jpeg;base64," + b64}}]}]}
        r = _post_json(url, headers, corpo)
        texto = ((r.get("choices") or [{}])[0].get("message") or {}).get("content") or ""
    if not texto.strip():
        raise RuntimeError("A IA devolveu uma resposta vazia.")
    return texto


def _txt(v) -> str:
    if isinstance(v, list):
        return "\n".join(str(x) for x in v).strip()
    return "" if v is None else str(v).strip()


def ler_com_ia(bgr, cfg: dict, dica: str = "") -> dict:
    texto = chamar_ia(cfg, foto_para_base64(bgr), PROMPT + (f"\n\nATENÇÃO: {dica}" if dica else ""))
    i, j = texto.find("{"), texto.rfind("}")
    try:
        bruto = json.loads(texto[i:j + 1])
    except ValueError:  # resposta fora do formato: mostra o texto cru para o usuário
        return {"transcricao": texto.strip(), "_sem_json": True}
    chaves = ("numero_processo", "data", "hora", "ocorrido", "materia", "transcricao", "observacoes")
    return {k: _txt(bruto.get(k)) for k in chaves}


# --------------------------------------------------------------------------
# Regras de negócio (sem interface)
# --------------------------------------------------------------------------
CNJ_RE = re.compile(r"(\d{7})\s*[-–—]\s*(\d{2})\s*\.\s*(\d{4})\s*\.\s*(\d)\s*\.\s*(\d{2})\s*\.\s*(\d{4})")
CNJ_EXATO = re.compile(r"(\d{7})-(\d{2})\.(\d{4})\.(\d)\.(\d{2})\.(\d{4})")


def extrair_processo(texto: str) -> str:
    """Normaliza o nº do processo: formato CNJ, 20 dígitos colados ou 15 a 20 dígitos."""
    if not texto:
        return ""
    m = CNJ_RE.search(texto)
    if m:
        return "{}-{}.{}.{}.{}.{}".format(*m.groups())
    m = re.search(r"(?<!\d)\d{20}(?!\d)", texto)
    if m:
        d = m.group(0)
        return f"{d[:7]}-{d[7:9]}.{d[9:13]}.{d[13]}.{d[14:16]}.{d[16:20]}"
    m = re.search(r"\d{15,20}", texto)
    return m.group(0) if m else ""


def cnj_valido(numero: str):
    """True/False conforme o dígito verificador (módulo 97, Res. CNJ 65/2008); None se fora do padrão CNJ."""
    m = CNJ_EXATO.fullmatch(numero or "")
    if not m:
        return None
    nnnnnnn, dv, ano, j, tr, oooo = m.groups()
    return int(nnnnnnn + ano + j + tr + oooo + dv) % 97 == 1


def validar_processo(numero: str):
    """Devolve (nivel, mensagem) para mostrar ao usuário: nivel = 'ok' | 'aviso' | ''."""
    numero = (numero or "").strip()
    if not numero:
        return "", ""
    v = cnj_valido(numero)
    if v is True:
        return "ok", "✔ número válido"
    if v is False:
        return "aviso", "⚠ dígito verificador não confere"
    return "aviso", "⚠ fora do padrão CNJ"


def normalizar_data_hora(data: str, hora: str = "") -> str:
    """Devolve 'dd/mm/aaaa' ou 'dd/mm/aaaa HH:MM'. Ano de 2 dígitos vira 4 (18 -> 2018)."""
    data, hora = (data or "").strip(), (hora or "").strip()
    saida = ""
    m = re.search(r"(\d{1,2})\s*[/.\-]\s*(\d{1,2})\s*[/.\-]\s*(\d{2}|\d{4})(?!\d)", data)
    if m:
        a = int(m.group(3))
        if len(m.group(3)) == 2:
            a += 2000 if a <= datetime.now().year % 100 else 1900
        try:
            saida = datetime(a, int(m.group(2)), int(m.group(1))).strftime("%d/%m/%Y")
        except ValueError:
            return data  # data impossível: deixa como a IA leu para o usuário corrigir
    else:
        saida = data
    mh = re.search(r"(\d{1,2})\s*[:hH]\s*(\d{2})?", hora)
    if mh and int(mh.group(1)) < 24 and int(mh.group(2) or 0) < 60:
        h = f"{int(mh.group(1)):02d}:{int(mh.group(2) or 0):02d}"
        saida = f"{saida} {h}".strip()
    return saida


def parse_data_texto(s):
    for fmt, com_hora in (("%d/%m/%Y %H:%M", True), ("%d/%m/%Y", False), ("%Y-%m-%d %H:%M:%S", True)):
        try:
            return datetime.strptime(s, fmt), com_hora
        except (TypeError, ValueError):
            pass
    return None, False


def abrir_banco() -> sqlite3.Connection:
    """NomeLote vazio = item do lote em andamento (ainda não exportado)."""
    db = sqlite3.connect(DB_PATH)
    db.execute(
        """CREATE TABLE IF NOT EXISTS levantamento_midias (
            ID INTEGER PRIMARY KEY AUTOINCREMENT,
            Title TEXT NOT NULL,
            Materia TEXT,
            DataHora TEXT,
            Secao TEXT,
            Estado_Midia TEXT,
            Observacao TEXT,
            NomeLote TEXT)""")
    db.commit()
    return db


def caminho_livre(pasta: Path, nome: str) -> Path:
    base = re.sub(r'[\\/:*?"<>|\r\n\t]+', "_", nome).strip(" .")[:100] or "lote"
    p, n = pasta / f"{base}.xlsx", 2
    while p.exists():
        p, n = pasta / f"{base}_{n}.xlsx", n + 1
    return p


def exportar_excel(linhas, caminho: Path):
    """linhas: (ID, Processo, Matéria, DataHora, Seção, Estado, Obs, Lote)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Levantamento"
    ws.append(["ID", "Nº Processo", "Matéria", "Data/Hora (do CD)", "Seção/Tipo", "Estado da Mídia", "Observação", "Lote"])
    formatos = []
    for r in linhas:
        r = list(r)
        dt, com_hora = parse_data_texto(r[3])
        if dt:
            r[3] = dt
        formatos.append("dd/mm/yyyy hh:mm" if com_hora else "dd/mm/yyyy")
        ws.append(r)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3A5F")
    for i, col in enumerate(ws.columns, start=1):
        larg = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[get_column_letter(i)].width = min(max(larg + 2, 12), 60)
    for row, fmt in zip(ws.iter_rows(min_row=2), formatos):
        row[3].number_format = fmt
        for k in (4, 6, 7):
            row[k].alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    wb.save(caminho)


# --------------------------------------------------------------------------
# Interface
# --------------------------------------------------------------------------
class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("OCR Mídias Judiciais V6")
        self.geometry("1100x700")
        self.minsize(1000, 660)
        self.configure(bg=BG)

        self.cap = None
        self.frame = None
        self.ultima_foto = None
        self.fila = queue.Queue()
        self.ocr_ocupado = False
        self.db = abrir_banco()
        self.cfg = carregar_config()

        self.proc, self.dt, self.secao = tk.StringVar(), tk.StringVar(), tk.StringVar()
        self.materia = tk.StringVar(value="A Verificar")
        self.estado = tk.StringVar(value="Bom")
        self.cam_idx = tk.StringVar(value="0")
        self.contador = tk.StringVar()
        self.info_ia = tk.StringVar()
        self.valida = tk.StringVar()
        self.status = tk.StringVar(value="Pronto. Inicie a câmera e pressione Espaço para capturar.")

        self._montar_ui()
        self.proc.trace_add("write", self._ao_mudar_processo)
        self._atualizar_info_ia()
        self.atualizar_contador()

        self.bind("<space>", self._tecla_espaco)
        self.bind_all("<Control-s>", self.salvar)
        self.protocol("WM_DELETE_WINDOW", self.fechar)
        self.after(100, self.checar_fila)
        if not self.cfg["chave"]:
            self.after(400, lambda: self.abrir_config("Informe a chave de API da IA para começar."))

    # ---- construção da tela ------------------------------------------------
    def _rotulo(self, pai, texto, bg=PAINEL):
        return tk.Label(pai, text=texto, bg=bg, fg=SUB, font=(FONTE, 9, "bold"), anchor="w")

    def _entrada(self, pai, var, **kw):
        return tk.Entry(pai, textvariable=var, bg=CAMPO, fg=kw.pop("fg", TXT), insertbackground=TXT, relief="flat",
                        font=(FONTE, 11), highlightthickness=1, highlightbackground=BORDA, highlightcolor=CIANO, **kw)

    def _texto(self, pai, linhas):
        return tk.Text(pai, height=linhas, bg=CAMPO, fg=TXT, insertbackground=TXT, relief="flat", font=(FONTE, 10),
                       wrap="word", highlightthickness=1, highlightbackground=BORDA, highlightcolor=CIANO)

    def _botao(self, pai, texto, cmd, cor, **kw):
        return tk.Button(pai, text=texto, command=cmd, bg=cor, fg="white", activebackground=cor,
                         activeforeground="white", relief="flat", cursor="hand2", takefocus=0,
                         highlightthickness=0, font=(FONTE, 10, "bold"), **kw)

    def _radios(self, pai, var, opcoes):
        f = tk.Frame(pai, bg=PAINEL)
        for o in opcoes:
            tk.Radiobutton(f, text=o, variable=var, value=o, bg=PAINEL, fg=TXT, selectcolor=CAMPO,
                           activebackground=PAINEL, activeforeground=TXT, takefocus=0, highlightthickness=0,
                           font=(FONTE, 10)).pack(side="left", padx=(0, 8))
        return f

    def _montar_ui(self):
        topo = tk.Frame(self, bg="#0f1017")
        topo.pack(fill="x")
        tk.Label(topo, text="", bg="#0f1017", fg=TXT, font=(FONTE, 10, "bold")).pack(side="left", padx=14, pady=10)
        tk.Label(topo, text="OCR MÍDIAS JUDICIAIS V6", bg="#0f1017", fg="white", font=("Georgia", 15, "bold")).pack(side="left", padx=30)
        self._botao(topo, "⚙ Configurar IA", self.abrir_config, "#374151").pack(side="right", padx=12, ipadx=6)
        tk.Label(topo, textvariable=self.info_ia, bg="#0f1017", fg=SUB, font=(FONTE, 9)).pack(side="right")

        # rodapé e status primeiro: ficam sempre visíveis, mesmo em telas baixas
        tk.Label(self, textvariable=self.status, bg=BG, fg=SUB, font=(FONTE, 9), anchor="w", justify="left",
                 wraplength=1040).pack(side="bottom", fill="x", padx=12, pady=(0, 6))
        rodape = tk.Frame(self, bg=PAINEL)
        rodape.pack(side="bottom", fill="x", padx=10, pady=(0, 6))
        tk.Label(rodape, textvariable=self.contador, bg="#4b5563", fg="white", font=(FONTE, 10, "bold"),
                 padx=14, pady=10).pack(side="left", padx=10, pady=8)
        self._botao(rodape, "📊 Concluir Lote (Exportar Excel)", self.concluir_lote, VERDE).pack(side="right", padx=10, ipadx=8, ipady=6)
        self._botao(rodape, "💾 Próximo CD (Salvar Item)", self.salvar, AZUL).pack(side="right", padx=4, ipadx=8, ipady=6)

        corpo = tk.Frame(self, bg=BG)
        corpo.pack(fill="both", expand=True, padx=10, pady=10)

        # Coluna esquerda: câmera + leitura da IA
        esq = tk.Frame(corpo, bg=PAINEL)
        esq.pack(side="left", fill="both", expand=True, padx=(0, 8))
        barra = tk.Frame(esq, bg=PAINEL)
        barra.pack(fill="x", padx=10, pady=(10, 4))
        self.btn_cam = self._botao(barra, "📷 Iniciar Câmera", self.alternar_camera, "#374151")
        self.btn_cam.pack(side="left")
        tk.Label(barra, text="Índice da câmera:", bg=PAINEL, fg=SUB, font=(FONTE, 9)).pack(side="left", padx=(14, 4))
        tk.Spinbox(barra, from_=0, to=9, width=3, textvariable=self.cam_idx, bg=CAMPO, fg=TXT,
                   buttonbackground=PAINEL, relief="flat").pack(side="left")

        self.moldura = tk.Frame(esq, bg="black", width=320, height=240)
        self.moldura.pack(fill="both", expand=True, padx=10, pady=4)
        self.moldura.pack_propagate(False)
        self.lbl_video = tk.Label(self.moldura, bg="black")
        self.lbl_video.pack(expand=True)

        self.btn_cap = self._botao(esq, "📷 Capturar e Ler com IA (Espaço)", self.capturar, "#374151")
        self.btn_cap.pack(pady=4, ipadx=10, ipady=3)

        self._rotulo(esq, "Leitura Extraída da Mídia (o que a IA viu)").pack(fill="x", padx=10)
        self.txt_ocr = self._texto(esq, 5)
        self.txt_ocr.config(state="disabled")
        self.txt_ocr.pack(fill="x", padx=10, pady=(2, 10))

        # Coluna direita: formulário
        dir_ = tk.Frame(corpo, bg=PAINEL, width=430)
        dir_.pack(side="right", fill="y", before=esq)  # reserva a largura do formulário primeiro
        dir_.pack_propagate(False)
        px = dict(padx=12, anchor="w")

        self._rotulo(dir_, "Captura Recente  (clique para ampliar)").pack(pady=(10, 2), **px)
        mold_prev = tk.Frame(dir_, bg="black", width=PREV_W, height=PREV_H)
        mold_prev.pack(padx=12)
        mold_prev.pack_propagate(False)
        self.lbl_prev = tk.Label(mold_prev, bg="black", cursor="hand2")
        self.lbl_prev.pack(expand=True)
        self.lbl_prev.bind("<Button-1>", self.ampliar_captura)

        lin = tk.Frame(dir_, bg=PAINEL)
        lin.pack(fill="x", padx=12, pady=(5, 1))
        self._rotulo(lin, "Nº Processo").pack(side="left")
        self.lbl_valida = tk.Label(lin, textvariable=self.valida, bg=PAINEL, fg=SUB, font=(FONTE, 9, "bold"))
        self.lbl_valida.pack(side="right")
        self.ent_proc = self._entrada(dir_, self.proc)
        self.ent_proc.pack(fill="x", padx=12, ipady=5)

        self._rotulo(dir_, "Matéria/Ramo").pack(pady=(5, 1), **px)
        self._radios(dir_, self.materia, ["Criminal", "Cível", "A Verificar"]).pack(**px)
        self._rotulo(dir_, "Data/Hora (do CD)").pack(pady=(5, 1), **px)
        self._entrada(dir_, self.dt, fg=CIANO).pack(fill="x", padx=12, ipady=5)
        self._rotulo(dir_, "Seção/Tipo (o que ocorreu)").pack(pady=(5, 1), **px)
        self._entrada(dir_, self.secao).pack(fill="x", padx=12, ipady=5)
        self._rotulo(dir_, "Estado da Mídia").pack(pady=(5, 1), **px)
        self._radios(dir_, self.estado, ["Bom", "Regular", "Ruim", "Não Funciona"]).pack(**px)
        self._rotulo(dir_, "Obs.").pack(pady=(5, 1), **px)
        self.txt_obs = self._texto(dir_, 3)
        self.txt_obs.pack(fill="x", padx=12, pady=(0, 10))

    def _set_leitura(self, texto):
        self.txt_ocr.config(state="normal")
        self.txt_ocr.delete("1.0", "end")
        self.txt_ocr.insert("1.0", texto)
        self.txt_ocr.config(state="disabled")

    def _ao_mudar_processo(self, *_):
        nivel, msg = validar_processo(self.proc.get())
        cor = {"ok": OK_COR, "aviso": AVISO_COR}.get(nivel, SUB)
        self.valida.set(msg)
        self.lbl_valida.config(fg=cor)
        self.ent_proc.config(highlightbackground=cor if nivel else BORDA)

    def _atualizar_info_ia(self):
        self.info_ia.set(f"IA: {PROVEDORES[self.cfg['provedor']]['nome']} · {self.cfg['modelo']}")

    # ---- configuração da IA ------------------------------------------------
    def abrir_config(self, aviso=None):
        if getattr(self, "_dlg", None) is not None and self._dlg.winfo_exists():
            self._dlg.lift()
            return
        dlg = self._dlg = tk.Toplevel(self)
        dlg.title("Configurar IA")
        dlg.configure(bg=PAINEL)
        dlg.resizable(False, False)
        dlg.transient(self)
        nomes = {v["nome"]: k for k, v in PROVEDORES.items()}
        nome_var = tk.StringVar(value=PROVEDORES[self.cfg["provedor"]]["nome"])
        chave, modelo = tk.StringVar(value=self.cfg["chave"]), tk.StringVar(value=self.cfg["modelo"])
        padrao = {p["modelo"] for p in PROVEDORES.values()}

        def trocar(_=None):  # se o modelo ainda é um padrão, acompanha o provedor escolhido
            if not modelo.get().strip() or modelo.get().strip() in padrao:
                modelo.set(PROVEDORES[nomes[nome_var.get()]]["modelo"])

        def gravar():
            k = chave.get().strip()
            if not k or not k.isascii() or " " in k:
                messagebox.showerror("Chave inválida", "Cole a chave de API completa, sem espaços ou aspas.", parent=dlg)
                return
            prov = nomes[nome_var.get()]
            self.cfg = {"provedor": prov, "chave": k, "modelo": modelo.get().strip() or PROVEDORES[prov]["modelo"]}
            salvar_config(self.cfg)
            self._atualizar_info_ia()
            self.status.set("Configuração da IA salva.")
            dlg.destroy()

        if aviso:
            tk.Label(dlg, text=aviso, bg=PAINEL, fg=AVISO_COR, font=(FONTE, 10, "bold")).pack(padx=16, pady=(14, 0), anchor="w")
        tk.Label(dlg, text="Provedor da IA", bg=PAINEL, fg=SUB, font=(FONTE, 9, "bold")).pack(padx=16, pady=(12, 2), anchor="w")
        om = tk.OptionMenu(dlg, nome_var, *nomes, command=trocar)
        om.config(bg=CAMPO, fg=TXT, activebackground=CAMPO, activeforeground=TXT, highlightthickness=0, relief="flat", width=30)
        om.pack(padx=16, anchor="w")
        tk.Label(dlg, text="Chave de API", bg=PAINEL, fg=SUB, font=(FONTE, 9, "bold")).pack(padx=16, pady=(10, 2), anchor="w")
        tk.Entry(dlg, textvariable=chave, show="•", width=52, bg=CAMPO, fg=TXT, insertbackground=TXT, relief="flat",
                 font=(FONTE, 10)).pack(padx=16, ipady=4)
        tk.Label(dlg, text="Modelo (troque aqui se o provedor desativar o modelo atual)", bg=PAINEL, fg=SUB,
                 font=(FONTE, 9, "bold")).pack(padx=16, pady=(10, 2), anchor="w")
        tk.Entry(dlg, textvariable=modelo, width=52, bg=CAMPO, fg=TXT, insertbackground=TXT, relief="flat",
                 font=(FONTE, 10)).pack(padx=16, ipady=4)
        tk.Label(dlg, text="A chave fica salva em config.ini, nesta pasta. Não compartilhe esse arquivo.", bg=PAINEL,
                 fg=SUB, font=(FONTE, 8)).pack(padx=16, pady=(8, 0), anchor="w")
        self._botao(dlg, "Salvar", gravar, AZUL).pack(pady=14, ipadx=24, ipady=3)
        try:
            dlg.grab_set()
        except tk.TclError:
            pass

    # ---- câmera ------------------------------------------------------------
    def alternar_camera(self):
        if self.cap:
            self.parar_camera()
            return
        try:
            idx = int(self.cam_idx.get())
        except ValueError:
            idx = 0
        cap = cv2.VideoCapture(idx, cv2.CAP_DSHOW) if os.name == "nt" else cv2.VideoCapture(idx)
        if not cap.isOpened():
            cap.release()
            messagebox.showerror("Câmera", f"Não foi possível abrir a câmera {idx}.\n"
                                 "Tente outro índice (1, 2...) ou feche outros programas que usem a câmera.")
            return
        cap.set(cv2.CAP_PROP_FRAME_WIDTH, 1920)
        cap.set(cv2.CAP_PROP_FRAME_HEIGHT, 1080)
        self.cap = cap
        self.btn_cam.config(text="⏹ Parar Câmera")
        self.status.set("Câmera ativa. Posicione a mídia e pressione Espaço.")
        self.loop_video()

    def parar_camera(self):
        if self.cap:
            self.cap.release()
        self.cap, self.frame = None, None
        self.lbl_video.config(image="")
        self.btn_cam.config(text="📷 Iniciar Câmera")

    def loop_video(self):
        if not self.cap:
            return
        ok, frame = self.cap.read()
        if ok:
            self.frame = frame
            self._mostrar(self.lbl_video, frame, max(self.moldura.winfo_width() - 4, 160),
                          max(self.moldura.winfo_height() - 4, 120))
        self.after(30, self.loop_video)

    @staticmethod
    def _mostrar(label, bgr, max_w, max_h):
        h, w = bgr.shape[:2]
        esc = min(max_w / w, max_h / h)
        pequena = cv2.resize(bgr, (max(1, int(w * esc)), max(1, int(h * esc))))
        img = ImageTk.PhotoImage(Image.fromarray(cv2.cvtColor(pequena, cv2.COLOR_BGR2RGB)))
        label.config(image=img)
        label.image = img  # mantém referência (senão o Tk apaga a imagem)

    def ampliar_captura(self, _=None):
        if self.ultima_foto is None:
            return
        top = tk.Toplevel(self)
        top.title("Captura ampliada (clique ou Esc para fechar)")
        top.configure(bg="black")
        lbl = tk.Label(top, bg="black", cursor="hand2")
        lbl.pack()
        self._mostrar(lbl, self.ultima_foto, int(self.winfo_screenwidth() * 0.85), int(self.winfo_screenheight() * 0.8))
        lbl.bind("<Button-1>", lambda e: top.destroy())
        top.bind("<Escape>", lambda e: top.destroy())
        top.focus_set()

    # ---- captura + leitura por IA ------------------------------------------
    def _tecla_espaco(self, _):
        foco = self.focus_get()
        if foco is not None and foco.winfo_class() in ("Entry", "Text", "Spinbox"):
            return  # o usuário está digitando
        self.capturar()
        return "break"

    def capturar(self):
        if self.frame is None:
            self.status.set("Inicie a câmera antes de capturar.")
            return
        if self.ocr_ocupado:
            return
        if not self.cfg["chave"]:
            self.abrir_config("Informe a chave de API da IA para ler as mídias.")
            return
        foto = self.frame.copy()
        self.limpar()  # evita sobrar dado do disco anterior
        self.ultima_foto = foto
        self._mostrar(self.lbl_prev, foto, PREV_W, PREV_H)
        self.focus_set()
        self.ocr_ocupado = True
        self.btn_cap.config(state="disabled")
        self.status.set("Lendo a mídia com IA... aguarde alguns segundos.")
        threading.Thread(target=self._worker, args=(foto, dict(self.cfg)), daemon=True).start()

    def _worker(self, foto, cfg):
        try:
            dados = ler_com_ia(foto, cfg)
            proc = extrair_processo(dados.get("numero_processo", ""))
            if proc and cnj_valido(proc) is False:  # dígito não confere: pede uma segunda leitura
                try:
                    dica = (f"Uma leitura anterior resultou em {proc}, cujo dígito verificador NÃO confere; "
                            "algum dígito provavelmente foi lido errado. Releia os dígitos um a um, com muito cuidado.")
                    d2 = ler_com_ia(foto, cfg, dica)
                    p2 = extrair_processo(d2.get("numero_processo", ""))
                    if p2 and cnj_valido(p2):
                        dados = d2
                except Exception:
                    pass
            self.fila.put(("ok", dados))
        except Exception as e:  # nunca deixa a thread morrer em silêncio
            self.fila.put(("erro", str(e) or e.__class__.__name__))

    def checar_fila(self):
        try:
            while True:
                tipo, d = self.fila.get_nowait()
                self.ocr_ocupado = False
                self.btn_cap.config(state="normal")
                if tipo == "erro":
                    self.status.set(f"A leitura por IA falhou: {d} Preencha manualmente ou capture de novo.")
                else:
                    self._aplicar_leitura(d)
        except queue.Empty:
            pass
        self.after(100, self.checar_fila)

    def _aplicar_leitura(self, d):
        bruto = d.get("numero_processo", "")
        proc = extrair_processo(bruto) or bruto
        data = normalizar_data_hora(d.get("data", ""), d.get("hora", ""))
        self.proc.set(proc)
        self.dt.set(data)
        self.secao.set(d.get("ocorrido", ""))
        self.materia.set({"criminal": "Criminal", "cível": "Cível", "civel": "Cível"}.get(
            d.get("materia", "").strip().lower(), "A Verificar"))
        self.txt_obs.insert("1.0", d.get("observacoes", ""))
        linhas = [f"Processo: {proc or '—'}", f"Data: {data or '—'}", f"Ocorrido: {d.get('ocorrido') or '—'}"]
        if d.get("transcricao"):
            linhas.append("Texto lido: " + " / ".join(x.strip() for x in d["transcricao"].splitlines() if x.strip()))
        self._set_leitura("\n".join(linhas))
        avisos = [validar_processo(proc)[1]] if proc else ["Nº do processo não identificado."]
        if not data:
            avisos.append("Sem data legível no disco.")
        if d.get("_sem_json"):
            avisos.append("A IA respondeu fora do formato esperado; confira o texto lido.")
        self.status.set("  ".join(a for a in avisos if a) + "  Confira e clique em Próximo CD (Ctrl+S).")

    # ---- salvar / lote -----------------------------------------------------
    def atualizar_contador(self):
        n = self.db.execute("SELECT COUNT(*) FROM levantamento_midias WHERE COALESCE(NomeLote,'')=''").fetchone()[0]
        self.contador.set(f"Mídias no Lote Atual: {n}")

    def limpar(self):
        self.proc.set("")
        self.dt.set("")
        self.secao.set("")
        self.materia.set("A Verificar")
        self.estado.set("Bom")
        self._set_leitura("")
        self.txt_obs.delete("1.0", "end")
        self.lbl_prev.config(image="")
        self.lbl_prev.image = None
        self.ultima_foto = None

    def salvar(self, _=None):
        if self.ocr_ocupado:
            self.status.set("Aguarde a leitura da IA terminar antes de salvar.")
            return "break"
        obs = self.txt_obs.get("1.0", "end").strip()
        proc, dt, secao = self.proc.get().strip(), self.dt.get().strip(), self.secao.get().strip()
        if not (proc or dt or secao or obs or self.ultima_foto is not None):
            self.status.set("Nada para salvar: capture uma mídia primeiro.")
            return "break"
        registro = (proc or "Sem Processo", self.materia.get() or "A Verificar", dt, secao,
                    self.estado.get() or "Bom", obs, "")
        try:
            with self.db:
                self.db.execute(
                    "INSERT INTO levantamento_midias "
                    "(Title, Materia, DataHora, Secao, Estado_Midia, Observacao, NomeLote) VALUES (?,?,?,?,?,?,?)",
                    registro)
        except sqlite3.Error as e:
            messagebox.showerror("Erro ao salvar", f"O item NÃO foi gravado:\n{e}")
            return "break"
        self.limpar()
        self.atualizar_contador()
        self.status.set(f"Item salvo ({registro[0]}). Posicione o próximo CD e pressione Espaço.")
        self.focus_set()
        return "break"

    def concluir_lote(self):
        if self.ocr_ocupado:
            self.status.set("Aguarde a leitura da IA terminar antes de exportar.")
            return
        linhas = self.db.execute(
            "SELECT ID, Title, Materia, DataHora, Secao, Estado_Midia, Observacao FROM levantamento_midias "
            "WHERE COALESCE(NomeLote,'')='' ORDER BY ID").fetchall()
        if not linhas:
            messagebox.showinfo("Lote vazio", "Não há itens salvos no lote atual.")
            return
        nome = simpledialog.askstring(
            "Concluir lote", f"{len(linhas)} mídias neste lote.\n\nNome do lote / informações do envelope:", parent=self)
        if nome is None:
            return
        nome = " ".join(nome.split())
        if not nome:
            messagebox.showwarning("Concluir lote", "Informe o nome do lote para exportar.")
            return
        EXPORT_DIR.mkdir(exist_ok=True)
        caminho = caminho_livre(EXPORT_DIR, nome)
        try:
            exportar_excel([r + (nome,) for r in linhas], caminho)
        except OSError as e:
            messagebox.showerror("Exportação", f"Não foi possível gravar o Excel (feche-o se estiver aberto):\n{e}")
            return
        with self.db:  # só marca como exportado depois que o Excel foi gravado
            self.db.executemany("UPDATE levantamento_midias SET NomeLote=? WHERE ID=?", [(nome, r[0]) for r in linhas])
        self.atualizar_contador()
        self.status.set(f"Lote '{nome}' exportado ({len(linhas)} mídias). Novo lote iniciado.")
        messagebox.showinfo("Lote exportado", f"{len(linhas)} mídias exportadas para:\n{caminho}\n\nUm novo lote foi iniciado.")
        if os.name == "nt":
            os.startfile(caminho)

    def fechar(self):
        self.parar_camera()
        self.db.close()
        self.destroy()


if __name__ == "__main__":
    if os.name == "nt":  # evita tela borrada em monitores com escala > 100%
        try:
            import ctypes
            ctypes.windll.shcore.SetProcessDpiAwareness(1)
        except Exception:
            pass
    App().mainloop()
